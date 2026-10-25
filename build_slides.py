"""Fills the duplicated LLM_AAT_template.pptx slides with the project content."""
import copy
import json
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
R = json.load(open("results/results.json"))
E1 = R["E1_models"]
TOPIC = "Sentiment Analysis of Movie Reviews: Lexicon, Classical ML and Transformer Embeddings"

RED = RGBColor(0xC0, 0x00, 0x00)
NAVY = RGBColor(0x1F, 0x3A, 0x68)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x59, 0x59, 0x59)
TINT = RGBColor(0xF3, 0xF5, 0xF9)
TINT_R = RGBColor(0xFB, 0xEE, 0xEE)
LINE = RGBColor(0xD0, 0xD5, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Times New Roman"

prs = Presentation(SRC)
slides = list(prs.slides)


def pct(x):
    return f"{100 * x:.1f}%"


# ---------------- title slide ----------------
for sh in slides[0].shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("TOPIC"):
        r0 = sh.text_frame.paragraphs[0].runs[0]
        r0.text = TOPIC
        r0.font.size = Pt(20)
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r.text = ""


# ---------------- helpers ----------------
def set_title(slide, text):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Agenda":
            sh.text_frame.paragraphs[0].runs[0].text = text
    for sh in list(slide.shapes):
        if sh.name == "Rectangle 2":
            sh._element.getparent().remove(sh._element)


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullet=False, space_after=4, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(paras, str):
        paras = [paras]
    for i, p_ in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        segs = p_.split("**")
        for k, seg in enumerate(segs):
            if not seg:
                continue
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = FONT
            f.size = Pt(size)
            f.bold = bold or k % 2 == 1
            f.color.rgb = color
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.2)))
            pPr.set("indent", str(-Inches(0.18)))
            from lxml import etree
            ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
            bf = etree.SubElement(pPr, f"{{{ns}}}buFont"); bf.set("typeface", "Arial")
            bc = etree.SubElement(pPr, f"{{{ns}}}buChar"); bc.set("char", "•")
    return tb


def card(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def stat(slide, x, y, w, value, label, color=RED):
    card(slide, x, y, w, 1.05)
    text(slide, x, y + 0.08, w, 0.55, value, size=26, color=color, bold=True, align=PP_ALIGN.CENTER)
    text(slide, x + 0.05, y + 0.62, w - 0.1, 0.4, label, size=11, color=MUTED, align=PP_ALIGN.CENTER)


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def table(slide, x, y, w, rows, col_w, size=11, row_h=0.3, header_fill=NAVY):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
    for j, cw in enumerate(col_w):
        gt.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        gt.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = gt.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (WHITE if i % 2 else TINT)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            segs = str(val).split("**")
            for k, seg in enumerate(segs):
                if not seg:
                    continue
                r = p.add_run(); r.text = seg
                r.font.name = FONT; r.font.size = Pt(size)
                r.font.bold = i == 0 or k % 2 == 1
                r.font.color.rgb = WHITE if i == 0 else INK
    return gt


def bar_chart(slide, x, y, w, h, cats, series, colors, title, minv=0.0, maxv=1.0, horizontal=False,
              fmt='0.00', legend=True):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    ct = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size = Pt(12); tp.runs[0].font.bold = True; tp.runs[0].font.name = FONT
    tp.runs[0].font.color.rgb = INK
    ch.font.name = FONT; ch.font.size = Pt(10); ch.font.color.rgb = MUTED
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = minv, maxv
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    va.tick_labels.number_format = fmt; va.tick_labels.number_format_is_linked = False
    ca = ch.category_axis
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(10)
    if horizontal:
        ca.reverse_order = True
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        va.tick_label_position = XL_TICK_LABEL_POSITION.HIGH  # keeps value axis at the bottom
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt; dl.number_format_is_linked = False
    dl.font.size = Pt(9); dl.font.color.rgb = INK
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for s, c in zip(plot.series, colors):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = c
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
    return ch


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t



# ===================================== SLIDES =====================================
D, E2, E3, E5, E6, MC = R["data"], R["E2_preprocessing"], R["E3_learning_curve"], R["E5_slices"], R["E6_negation_posthoc"], R["E4_mcnemar"]
VAD, NB, LR_, SVM, WV, RB, ENS = ["VADER (lexicon, rule-based)", "Multinomial Naive Bayes (BoW)",
                                  "Logistic Regression (TF-IDF)", "Linear SVM (TF-IDF)", "Static word vectors + LR",
                                  "RoBERTa embeddings (frozen) + LR", "Ensemble (TF-IDF LR + RoBERTa LR)"]
S = slides
BLUE, ORANGE = RGBColor(0x1F, 0x3A, 0x68), RGBColor(0xC0, 0x00, 0x00)

# 3 Introduction
s = S[2]; set_title(s, "Introduction")
card(s, 0.45, 0.8, 4.3, 4.45, fill=TINT)
text(s, 0.6, 0.9, 4.0, 0.4, "Why sentiment analysis?", size=18, color=NAVY, bold=True)
text(s, 0.6, 1.35, 4.0, 3.8, [
    "Companies, studios and platforms track millions of opinions in reviews and social media.",
    "Short opinion sentences are hard: **negation** ('not good'), **contrast** ('stunning, but dull'), "
    "**sarcasm**.",
    "NLP offers three generations of solutions: **lexicons**, **statistical ML** and **Transformer** "
    "embeddings.",
    "Which works best, and why? We measure them side by side."], size=14, bullet=True, space_after=6)
text(s, 5.0, 0.85, 4.5, 0.4, "Objectives", size=18, color=RED, bold=True)
objs = ["Identify the task, a real-time application and APIs",
        "Measure the effect of Unit 1 preprocessing",
        "Compare 7 approaches: lexicon, NB, LR, SVM, word vectors, RoBERTa, ensemble",
        "Evaluate with accuracy, P, R, F1, ROC-AUC + significance test",
        "Analyse negation, contrast, sarcasm and data efficiency",
        "Deliver a working, explainable sentiment analyser"]
for i, o in enumerate(objs):
    yy = 1.3 + i * 0.66
    card(s, 5.0, yy, 0.42, 0.42, fill=RED, shape=MSO_SHAPE.OVAL)
    text(s, 5.0, yy + 0.04, 0.42, 0.34, str(i + 1), size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, 5.52, yy - 0.02, 4.0, 0.6, o, size=13, anchor=MSO_ANCHOR.MIDDLE)

# 4 Topic overview
s = S[3]; set_title(s, "Selected NLP Topic Overview")
card(s, 0.45, 0.8, 9.1, 1.0, fill=TINT)
text(s, 0.6, 0.85, 8.8, 0.95, [
    "**Task:** binary text classification — given a review sentence x, predict y ∈ {negative, positive}.",
    "**Classifier:** learn f(x) = P(positive | x) from labelled examples; predict positive if f(x) ≥ 0.5. "
    "**Lexicon:** sum word valences and apply rules."], size=13, space_after=3)
text(s, 0.45, 1.95, 4.4, 0.4, "Real-world use cases", size=16, color=RED, bold=True)
text(s, 0.45, 2.35, 4.4, 2.9, ["**Product reviews:** flag unhappy customers",
                              "**Brand monitoring** on social media", "**Film reviews:** audience reaction",
                              "**Support tickets:** prioritise angry customers",
                              "**Finance:** news sentiment as a market signal"], size=13, bullet=True, space_after=5)
table(s, 5.05, 2.0, 4.5, [
    ["Unit", "Used in this project"],
    ["1", "Regex, normalisation, tokenisation, stopwords, negation"],
    ["2", "Bag-of-words, TF-IDF, static and contextual embeddings"],
    ["3", "N-grams, RoBERTa Transformer, self-attention"],
    ["4", "Case study: NLP-powered opinion mining"],
    ["5", "Deployment trade-offs, responsible AI"],
], [0.55, 3.95], size=11, row_h=0.47)

# 5 Real-time application
s = S[4]; set_title(s, "Real-Time Application / API Identified")
text(s, 0.45, 0.75, 9.1, 0.65, "**Movie-Review Sentiment Analyser** (predict.py) built on **NLTK VADER** (rule-based "
     "API), **scikit-learn** and **RoBERTa-base** via spaCy. Cloud equivalents: Google Cloud Natural Language, "
     "Amazon Comprehend, Azure AI Language.", size=13)
tc0 = R["test_cases"][0]
stages = [("INPUT", "One review sentence", f"\"{tc0['text']}\""),
          ("PROCESSING", "Clean → tokenise → TF-IDF + RoBERTa → 2 × Logistic Regression → average",
           f"TF-IDF LR {tc0['tfidf_lr_p_pos']:.2f}, RoBERTa LR {tc0['roberta_lr_p_pos']:.2f}"),
          ("OUTPUT", "Label, P(positive), key words, VADER baseline", "POSITIVE (0.99); masterpiece (+)")]
for i, (h, d, ex) in enumerate(stages):
    xx = 0.45 + i * 3.1
    card(s, xx, 1.5, 2.85, 2.3, fill=TINT if i != 1 else TINT_R)
    text(s, xx + 0.1, 1.58, 2.65, 0.35, h, size=15, color=RED if i == 1 else NAVY, bold=True)
    text(s, xx + 0.1, 1.95, 2.65, 0.8, d, size=13, bold=True)
    text(s, xx + 0.1, 2.85, 2.65, 0.9, ex, size=11, color=MUTED)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(xx + 2.88), Inches(2.5), Inches(0.2), Inches(0.3))
        a.fill.solid(); a.fill.fore_color.rgb = MUTED; a.line.fill.background()
text(s, 0.45, 3.95, 4.5, 0.35, "Functionality", size=14, color=NAVY, bold=True)
text(s, 0.45, 4.3, 4.5, 1.1, ["Label + confidence + explanation", "Fast mode (TF-IDF, ~0.02 ms)",
                             "Accurate mode (ensemble with RoBERTa)"], size=12, bullet=True, space_after=2)
text(s, 5.05, 3.95, 4.5, 0.35, "Limitations", size=14, color=RED, bold=True)
text(s, 5.05, 4.3, 4.5, 1.1, ["Trained on critic snippets only", "Binary: no neutral or mixed class",
                             "Sarcasm not handled reliably"], size=12, bullet=True, space_after=2)

# 6 Architecture
s = S[5]; set_title(s, "NLP Model / Architecture Study")
picture(s, "figures/fig_architecture.png", 0.35, 0.75, w=6.2)
text(s, 6.7, 0.75, 2.9, 0.35, "6 components", size=16, color=RED, bold=True)
text(s, 6.7, 1.12, 2.9, 3.0, ["**Input** sentence", "**Preprocessing** (clean, tokenise, negation)",
                             "**Feature extractors** (TF-IDF, vectors, RoBERTa, lexicon)",
                             "**Classifiers** (NB, LR, SVM, rule)", "**Ensemble** (average)",
                             "**Output** label + explanation"], size=13, bullet=True, space_after=4)
card(s, 0.45, 4.2, 9.1, 1.15, fill=TINT)
text(s, 0.6, 4.25, 8.8, 1.05, ["**RoBERTa-base:** 12 Transformer layers × 12-head self-attention, "
                              "softmax(QKᵀ/√d)V, 768-d token states, byte-level BPE (50,265 tokens). Sentence "
                              "vector = mean of last-layer states (frozen, no fine-tuning)."], size=12)

# 7 Models compared
s = S[6]; set_title(s, "Approaches and Language Models Compared")
table(s, 0.45, 0.8, 5.6, [
    ["#", "Approach", "Features → decision"],
    ["1", "VADER", "Valence lexicon + rules → compound ≥ 0.05"],
    ["2", "Multinomial Naive Bayes", "Word counts → Bayes' rule"],
    ["3", "Logistic Regression", "TF-IDF → σ(wᵀx + b)"],
    ["4", "Linear SVM", "TF-IDF → max-margin hyperplane"],
    ["5", "Static word vectors + LR", "Mean 300-d vectors → LR"],
    ["6", "RoBERTa + LR", "Mean 768-d contextual → LR"],
    ["7", "Ensemble", "Average of LR and RoBERTa-LR"],
], [0.35, 2.05, 3.2], size=11, row_h=0.45)
card(s, 6.3, 0.8, 3.25, 4.5, fill=TINT)
text(s, 6.45, 0.88, 3.0, 0.35, "Key formulas", size=15, color=RED, bold=True)
text(s, 6.45, 1.28, 3.0, 4.0, ["**TF-IDF:** (1 + log tf) · idf",
                              "**NB:** argmax log P(c) + Σ log P(t | c)",
                              "**LR:** P(+|x) = 1 / (1 + e^−(wᵀx+b))",
                              "**SVM:** min ½‖w‖² + C Σ hinge",
                              "**F1:** 2PR / (P + R)",
                              "**ROC-AUC:** area under TPR vs FPR"], size=12, space_after=8)

# 8 Dataset & preprocessing
s = S[7]; set_title(s, "Dataset and Preprocessing")
for i, (v, l) in enumerate([(f"{D['n']:,}", "review sentences (Pang & Lee 2005)"), ("50 / 50", "positive / negative"),
                            (f"{D['avg_words']:.0f}", "words per sentence (avg.)"),
                            (f"{D['n_train']:,} / {D['n_test']:,}", "train / test (stratified)")]):
    stat(s, 0.45 + i * 2.32, 0.8, 2.12, v, l, color=RED if i % 2 == 0 else NAVY)
text(s, 0.45, 2.0, 4.1, 3.4, [
    "**Source:** Rotten Tomatoes snippets, sentence polarity v1.0",
    "**Cleaning:** UTF-8, HTML remnants, whitespace, lowercase",
    "**Tokenisation:** PTB-style tokens (n't, 's split); BPE for RoBERTa",
    "**Stopword removal hurts:** removes 'not', 'too', 'very'",
    "**Negation marking:** not a good film → not NOT_a NOT_good NOT_film",
    "Choices made by **5-fold CV** on train only"], size=12, bullet=True, space_after=4)
ks = list(E2)
bar_chart(s, 4.7, 1.95, 4.9, 3.45, ["Unigrams", "− stopwords", "Uni+bigrams", "+ negation"],
          [("5-fold CV accuracy", [E2[k]["cv_acc_mean"] for k in ks])], [BLUE], "Preprocessing ablation (TF-IDF + LR)",
          minv=0.70, maxv=0.80, fmt='0.000', legend=False)

# 9 Implementation
s = S[8]; set_title(s, "Implementation Methodology")
steps = ["Load corpus, stratified 80/20 split", "Preprocessing ablation (5-fold CV)", "VADER on test set (no training)",
         "Grid-search NB, LR, SVM", "Extract word-vector and RoBERTa embeddings (cached)",
         "Train LR on embeddings; build ensemble", "Metrics, learning curve, McNemar, slices", "Package predict.py"]
for i, st in enumerate(steps):
    yy = 0.8 + i * 0.57
    card(s, 0.45, yy, 3.0, 0.47, fill=TINT if i < 4 else TINT_R)
    text(s, 0.55, yy + 0.02, 2.85, 0.43, f"{i + 1}. {st}", size=11, anchor=MSO_ANCHOR.MIDDLE)
bp = lambda n: ", ".join(f"{k.split('__')[-1]} = {v}" for k, v in E1[n]["best_params"].items())
table(s, 3.7, 0.8, 5.85, [
    ["Item", "Setting"],
    ["Environment", "Python 3.13, CPU only (2 cores, 8 GB)"],
    ["Libraries", "scikit-learn, NLTK (VADER), spaCy 3.8, PyTorch"],
    ["Naive Bayes", bp(NB)],
    ["Logistic Regression", bp(LR_)],
    ["Linear SVM", bp(SVM)],
    ["RoBERTa + LR", bp(RB) + " (standardised features)"],
    ["Ensemble", "Equal-weight average of probabilities"],
    ["Protocol", "Tune on train (5-fold CV); test used once"],
], [1.7, 4.15], size=11, row_h=0.47)

# 10 Results comparison
s = S[9]; set_title(s, "Results: Model Comparison")
order = [VAD, NB, LR_, SVM, WV, RB, ENS]
labs = ["VADER", "Naive Bayes", "Logistic Reg.", "Linear SVM", "Word vectors", "RoBERTa + LR", "Ensemble"]
bar_chart(s, 0.4, 0.75, 6.0, 4.65, labs, [("Accuracy", [E1[n]["accuracy"] for n in order]),
                                         ("F1", [E1[n]["f1"] for n in order])], [BLUE, ORANGE],
          f"Held-out test set (n = {D['n_test']:,})", minv=0.5, maxv=0.9, horizontal=True, fmt='0.000')
stat(s, 6.6, 0.8, 1.42, f"{E1[ENS]['accuracy']:.3f}", "Ensemble accuracy")
stat(s, 8.13, 0.8, 1.42, f"{E1[ENS]['roc_auc']:.3f}", "Ensemble ROC-AUC", color=NAVY)
text(s, 6.6, 2.0, 2.95, 3.4, [
    f"**Ensemble best:** +{100 * (E1[ENS]['accuracy'] - E1[SVM]['accuracy']):.1f} pts over best single model",
    f"**SVM ≈ LR ≈ NB** ({E1[SVM]['accuracy']:.3f} / {E1[LR_]['accuracy']:.3f} / {E1[NB]['accuracy']:.3f}); "
    f"McNemar p = {MC['p_value']:.2f}",
    f"**VADER** {E1[VAD]['accuracy']:.3f}: {100 * R['vader_neutral_fraction']:.0f}% scored neutral",
    f"**Frozen RoBERTa** {E1[RB]['accuracy']:.3f} alone"], size=12, bullet=True, space_after=5)

# 11 Learning curve + slices
s = S[10]; set_title(s, "Results: Data Efficiency and Linguistic Patterns")
picture(s, "figures/fig_learning_curve.png", 0.35, 0.8, w=4.6)
picture(s, "figures/fig_slices.png", 5.05, 0.8, w=4.6)
f0 = E3[sorted(E3, key=float)[0]]
card(s, 0.45, 3.55, 4.45, 1.2, fill=TINT)
text(s, 0.55, 3.62, 4.25, 1.1, [f"With only {f0['n']} training sentences RoBERTa leads ({f0['roberta_lr']:.3f} vs "
                               f"{f0['tfidf_lr']:.3f}); TF-IDF overtakes after ~4,000 sentences.",
                               "Pre-training = prior knowledge; frozen features saturate."], size=12, bullet=True,
     space_after=4)
card(s, 5.1, 3.55, 4.45, 1.2, fill=TINT_R)
text(s, 5.2, 3.62, 4.25, 1.1, [f"Negation: TF-IDF LR {E5['Contains negation'][LR_]:.3f}, VADER "
                              f"{E5['Contains negation'][VAD]:.3f}",
                              f"Negation marking (post-hoc) lifts this to {E6['negation_slice_acc']:.3f}",
                              f"Contrast hurts RoBERTa most ({E5[list(E5)[1]][RB]:.3f})"], size=12, bullet=True,
     space_after=4)

# 12 Test cases
s = S[11]; set_title(s, "Demonstrated Test Cases")


def mk(p, e):
    return ("✓ " if (p >= 0.5) == bool(e) else "✗ ") + f"{p:.2f}"


rows = [["#", "Sentence", "Gold", "TF-IDF LR", "RoBERTa", "LR+neg", "VADER"]]
for i, t in enumerate(R["test_cases"]):
    vd = t["vader_compound"]
    rows.append([f"TC{i + 1}", t["text"], "Pos" if t["expected"] else "Neg", mk(t["tfidf_lr_p_pos"], t["expected"]),
                 mk(t["roberta_lr_p_pos"], t["expected"]), mk(E6["test_cases"][t["text"]], t["expected"]),
                 ("✓ " if (vd >= 0.05) == bool(t["expected"]) else "✗ ") + f"{vd:+.2f}"])
table(s, 0.45, 0.78, 9.1, rows, [0.5, 4.2, 0.55, 1.0, 0.95, 0.95, 0.95], size=10, row_h=0.5)

# 13 Discussion
s = S[12]; set_title(s, "Discussion and Error Analysis")
text(s, 0.45, 0.78, 4.4, 0.35, "Key observations", size=16, color=NAVY, bold=True)
text(s, 0.45, 1.15, 4.4, 4.2, ["Simple linear models on TF-IDF remain hard to beat with enough data",
                              "Transformers help most with little data and in ensembles",
                              "Stopword removal hurts sentiment (negations are stopwords)",
                              "Averaged static vectors blur sentiment words",
                              "A generic lexicon (VADER) misses film vocabulary"], size=13, bullet=True,
     space_after=6)
errs = [("Contrast / concession", "'more depressing than liberating, but it's never boring' → wrong"),
        ("Double negation", "'isn't nearly as terrible as it could have been' → wrong"),
        ("Faint praise", "'a solid film … but more conscientious than stirring' → wrong"),
        ("Sarcasm / casual style", "'oh great, another two hours…' fools VADER")]
text(s, 5.05, 0.78, 4.5, 0.35, "Error types", size=16, color=RED, bold=True)
for i, (h, d) in enumerate(errs):
    yy = 1.2 + i * 1.03
    card(s, 5.05, yy, 4.5, 0.9, fill=TINT_R if i % 2 == 0 else TINT)
    text(s, 5.15, yy + 0.05, 4.3, 0.35, h, size=13, bold=True)
    text(s, 5.15, yy + 0.4, 4.3, 0.45, d, size=11, color=MUTED)

# 14 Conclusion
s = S[13]; set_title(s, "Conclusion and Future Work")
card(s, 0.45, 0.8, 4.4, 2.75, fill=TINT)
text(s, 0.6, 0.88, 4.1, 0.35, "Summary", size=16, color=NAVY, bold=True)
text(s, 0.6, 1.28, 4.1, 3.0, [f"7 approaches compared on {D['n']:,} sentences",
                             f"Best: ensemble {100 * E1[ENS]['accuracy']:.1f}% accuracy, AUC {E1[ENS]['roc_auc']:.3f}",
                             f"Best single: Linear SVM {100 * E1[SVM]['accuracy']:.1f}%",
                             f"VADER without training: {100 * E1[VAD]['accuracy']:.1f}%",
                             "Working, explainable analyser delivered"], size=13, bullet=True, space_after=6)
card(s, 5.05, 0.8, 4.5, 2.75, fill=TINT_R)
text(s, 5.2, 0.88, 4.2, 0.35, "Future work", size=16, color=RED, bold=True)
text(s, 5.2, 1.28, 4.2, 3.0, ["Fine-tune RoBERTa / DistilBERT (or LoRA) on GPU", "Negation scope + contrast-aware "
                             "weighting", "Zero/few-shot LLM prompting comparison", "Neutral class and aspect-based "
                             "sentiment", "Indian-language and code-mixed reviews"], size=13, bullet=True, space_after=6)
text(s, 0.45, 3.9, 9.1, 0.8, "Classical features + pre-trained Transformer knowledge together give the most "
     "reliable sentiment classifier.", size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 15 References
s = S[14]; set_title(s, "References")
refs = ["Jurafsky and Martin, Speech and Language Processing, 3rd ed., 2024 (Ch. 4)",
        "Pang and Lee, Seeing Stars, ACL 2005 (sentence polarity dataset)",
        "Pang, Lee and Vaithyanathan, Thumbs up? Sentiment Classification, EMNLP 2002",
        "Hutto and Gilbert, VADER, ICWSM 2014",
        "Das and Chen, negation marking, APFA 2001",
        "Wang and Manning, Baselines and Bigrams, ACL 2012",
        "Socher et al., Sentiment Treebank, EMNLP 2013",
        "Liu et al., RoBERTa, arXiv:1907.11692, 2019",
        "Vaswani et al., Attention Is All You Need, NeurIPS 2017",
        "McNemar, Psychometrika, 1947",
        "Tunstall et al., NLP with Transformers, O'Reilly, 2022"]
text(s, 0.45, 0.8, 9.1, 4.6, [f"[{i + 1}] {r}" for i, r in enumerate(refs)], size=13, space_after=4)

prs.save(OUT)
print("saved", OUT)
